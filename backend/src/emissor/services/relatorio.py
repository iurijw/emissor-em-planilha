"""Planilha das emissões filtradas: uma linha por nota, com ISS, retenções e líquido lidos do XML da NFS-e.

Os valores vêm do XML autorizado (o que a Sefin calculou), não da tabela. Notas canceladas ou substituídas aparecem
com a situação, em cinza.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from emissor.db import Emissao
from emissor.nfse import codes
from emissor.nfse.documentos import formatar_documento
from emissor.nfse.models import RETENCOES_PIS_COFINS
from emissor.nfse.xml_reader import NFSeDoc, NfseXmlError
from emissor.services import arquivos
from emissor.services.tabela import _texto_csv

log = logging.getLogger(__name__)

MAX_NOTAS = 20_000

SITUACAO = {"autorizada": "Autorizada", "cancelada": "Cancelada", "substituida": "Substituída"}
ZERO = Decimal("0")


@dataclass
class LinhaRelatorio:
    competencia: str  # MM/AAAA
    numero: str
    emissao: str  # dd/mm/aaaa
    situacao: str
    tomador: str
    documento: str
    municipio_tomador: str
    servico: str  # cTribNac
    descricao: str
    valor: Decimal
    base_iss: Decimal | None
    aliquota_iss: Decimal | None
    iss: Decimal | None
    iss_retido: bool
    irrf: Decimal
    csll: Decimal
    inss: Decimal
    pis: Decimal
    cofins: Decimal
    total_retido: Decimal
    liquido: Decimal
    tributacao: str
    ambiente: str
    chave: str
    conta: bool  # autorizada (canceladas e substituídas saem em cinza)
    ordem: str  # AAAAMM, para ordenar

    @property
    def iss_retido_valor(self) -> Decimal:
        return (self.iss or ZERO) if self.iss_retido else ZERO


def _dec(v: str | None) -> Decimal | None:
    if v in (None, ""):
        return None
    try:
        return Decimal(v)
    except InvalidOperation:
        return None


def _linha(e: Emissao, doc: NFSeDoc) -> LinhaRelatorio:
    d, v = doc.dps, doc.t
    tp_ret_iss = d("valores/trib/tribMun/tpRetISSQN")
    pc = "valores/trib/tribFed/piscofins"
    pis_ret, cofins_ret, _ = RETENCOES_PIS_COFINS.get(d(f"{pc}/tpRetPisCofins") or "", (False, False, None))
    irrf = _dec(d("valores/trib/tribFed/vRetIRRF")) or ZERO
    csll = _dec(d("valores/trib/tribFed/vRetCSLL")) or ZERO
    inss = _dec(d("valores/trib/tribFed/vRetCP")) or ZERO
    pis = (_dec(d(f"{pc}/vPis")) or ZERO) if pis_ret else ZERO
    cofins = (_dec(d(f"{pc}/vCofins")) or ZERO) if cofins_ret else ZERO
    valor = doc.valor_servico
    iss = _dec(v("valores/vISSQN"))
    iss_retido = tp_ret_iss in ("2", "3")
    total_ret = _dec(v("valores/vTotalRet"))
    if total_ret is None:
        total_ret = irrf + csll + inss + pis + cofins + ((iss or ZERO) if iss_retido else ZERO)
    liquido = _dec(v("valores/vLiq"))
    if liquido is None:
        liquido = valor - total_ret
    compet = d("dCompet") or ""  # AAAA-MM-DD
    dh = d("dhEmi") or ""
    c_mun = d("toma/end/endNac/cMun")
    return LinhaRelatorio(
        competencia=f"{compet[5:7]}/{compet[:4]}" if len(compet) >= 7 else "",
        numero=doc.numero or "",
        emissao=f"{dh[8:10]}/{dh[5:7]}/{dh[:4]}" if len(dh) >= 10 else "",
        situacao=SITUACAO.get(e.status, e.status),
        tomador=doc.tomador_nome or e.tomador_nome,
        documento=formatar_documento(doc.tomador_documento or e.tomador_documento),
        municipio_tomador=(codes.municipio_uf(c_mun) or c_mun) if c_mun else "",
        servico=d("serv/cServ/cTribNac") or "",
        descricao=d("serv/cServ/xDescServ") or e.descricao,
        valor=valor,
        base_iss=_dec(v("valores/vBC")),
        aliquota_iss=_dec(v("valores/pAliqAplic")),
        iss=iss,
        iss_retido=iss_retido,
        irrf=irrf,
        csll=csll,
        inss=inss,
        pis=pis,
        cofins=cofins,
        total_retido=total_ret,
        liquido=liquido,
        tributacao=e.tributacao_nome or ("Só desta nota" if e.tributacao_json else "Padrão"),
        ambiente="Produção" if e.ambiente == "1" else "Homologação",
        chave=e.chave_acesso or doc.chave,
        conta=e.status == "autorizada",
        ordem=compet[:4] + compet[5:7],
    )


def linhas(emissoes: list[Emissao]) -> list[LinhaRelatorio]:
    res = []
    for e in emissoes:
        p = arquivos.absoluto(e.xml_path)
        if not p or not p.is_file():
            log.warning("relatório: XML da nota não encontrado", extra={"emissao_id": e.id})
            continue
        try:
            res.append(_linha(e, NFSeDoc.from_bytes(p.read_bytes())))
        except (NfseXmlError, OSError):
            log.exception("relatório: XML da nota ilegível", extra={"emissao_id": e.id})
    res.sort(key=lambda r: (r.ordem, int(r.numero) if r.numero.isdigit() else 0))
    return res


COLUNAS = [
    ("Competência", 12),
    ("Nº NFS-e", 10),
    ("Emissão", 12),
    ("Situação", 12),
    ("Tomador", 40),
    ("CNPJ/CPF", 20),
    ("Município do tomador", 24),
    ("Código do serviço", 12),
    ("Descrição", 50),
    ("Valor do serviço", 15),
    ("Base do ISS", 15),
    ("Alíquota do ISS (%)", 12),
    ("ISS", 13),
    ("ISS retido pelo tomador", 13),
    ("IRRF retido", 13),
    ("CSLL retida", 13),
    ("INSS retido", 13),
    ("PIS retido", 13),
    ("COFINS retido", 13),
    ("Total retido", 13),
    ("Valor líquido", 15),
    ("Tributação", 18),
    ("Ambiente", 13),
    ("Chave de acesso", 54),
]
COLUNAS_MOEDA = {10, 11, 13, 14, 15, 16, 17, 18, 19, 20, 21}  # 1 = primeira coluna


def _valores(r: LinhaRelatorio) -> list:
    return [
        r.competencia,
        r.numero,
        r.emissao,
        r.situacao,
        r.tomador,
        r.documento,
        r.municipio_tomador,
        r.servico,
        r.descricao,
        r.valor,
        r.base_iss,
        r.aliquota_iss,
        r.iss,
        r.iss_retido_valor,
        r.irrf,
        r.csll,
        r.inss,
        r.pis,
        r.cofins,
        r.total_retido,
        r.liquido,
        r.tributacao,
        r.ambiente,
        r.chave,
    ]


def _cabecalho(ws, colunas: list[tuple[str, int]]) -> None:
    ws.append([c for c, _ in colunas])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0098D7")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, (_, larg) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.freeze_panes = "A2"


MOEDA = '"R$" #,##0.00'


def xlsx_bytes(rs: list[LinhaRelatorio]) -> bytes:
    wb = Workbook()
    wn = wb.active
    wn.title = "Notas"
    _cabecalho(wn, COLUNAS)
    for r in rs:
        wn.append(_valores(r))
        linha = wn.max_row
        for col in (5, 9):  # texto digitado nunca vira fórmula na planilha
            c = wn.cell(row=linha, column=col)
            if c.data_type == "f":
                c.data_type = "s"
        for col in COLUNAS_MOEDA:
            wn.cell(row=linha, column=col).number_format = MOEDA
        wn.cell(row=linha, column=12).number_format = "0.00"
        if not r.conta:
            for c in wn[linha]:
                c.font = Font(color="8A8F94")
    if rs:
        wn.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{wn.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _br(v) -> str:
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return f"{v:.2f}".replace(".", ",")
    return _texto_csv(v)


def csv_bytes(rs: list[LinhaRelatorio]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c for c, _ in COLUNAS])
    for r in rs:
        w.writerow([_br(x) for x in _valores(r)])
    return buf.getvalue().encode("utf-8-sig")
