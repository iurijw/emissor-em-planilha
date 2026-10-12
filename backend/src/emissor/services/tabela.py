"""Tabela de notas a emitir: normalização, importação e exportação (XLSX/CSV)."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from emissor.nfse.models import normalizar_documento

COLUNAS = ["Nome", "CNPJ/CPF", "Valor", "Descrição"]

_SINONIMOS = {
    "nome": ("nome", "empresa", "razao social", "razão social", "cliente", "tomador", "nome da empresa"),
    "documento": ("cnpj", "cpf", "cnpj/cpf", "cpf/cnpj", "documento", "cnpj cpf"),
    "valor": ("valor", "valor (r$)", "valor r$", "preco", "preço", "total", "honorarios", "honorários"),
    "descricao": ("descricao", "descrição", "servico", "serviço", "descricao do servico", "discriminacao"),
}


class ImportacaoError(ValueError):
    pass


class _PontoEVirgula(csv.excel):
    delimiter = ";"


@dataclass
class LinhaImportada:
    nome: str
    documento: str
    valor: str
    descricao: str


@dataclass
class ResultadoImportacao:
    linhas: list[LinhaImportada] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def parse_valor(v) -> str:
    """Aceita número, "1.234,56", "1234.56", "R$ 1.234,56". Retorna "1234.56" ou "" se vazio."""
    if v is None:
        return ""
    if isinstance(v, int | float | Decimal):
        return f"{Decimal(str(v)).quantize(Decimal('0.01')):.2f}"
    s = str(v).strip().replace("R$", "").replace(" ", "").replace(" ", "")
    if not s:
        return ""
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return f"{Decimal(s).quantize(Decimal('0.01')):.2f}"
    except InvalidOperation as exc:
        raise ValueError(f"valor inválido: {v!r}") from exc


def _mapear_cabecalho(cab: list) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for i, c in enumerate(cab):
        n = _norm(c)
        for campo, nomes in _SINONIMOS.items():
            if campo not in mapa and n in {_norm(x) for x in nomes}:
                mapa[campo] = i
    return mapa


def _linhas_para_resultado(linhas: list[list]) -> ResultadoImportacao:
    linhas = [row for row in linhas if any(str(c).strip() for c in row if c is not None)]
    if not linhas:
        raise ImportacaoError("O arquivo está vazio.")
    mapa = _mapear_cabecalho(linhas[0])
    if len(mapa) >= 2:
        corpo = linhas[1:]
    else:
        # Sem cabeçalho reconhecível: assume a ordem Nome, CNPJ, Valor, Descrição.
        mapa = {"nome": 0, "documento": 1, "valor": 2, "descricao": 3}
        corpo = linhas
    faltando = [c for c in ("documento", "valor") if c not in mapa]
    if faltando:
        raise ImportacaoError(
            "Não encontrei as colunas obrigatórias: "
            + ", ".join({"documento": "CNPJ/CPF", "valor": "Valor"}[c] for c in faltando)
            + f". Use os cabeçalhos: {', '.join(COLUNAS)}."
        )
    res = ResultadoImportacao()

    def cel(row: list, campo: str):
        i = mapa.get(campo)
        return row[i] if i is not None and i < len(row) else None

    for n, row in enumerate(corpo, start=2 if corpo is not linhas else 1):
        try:
            valor = parse_valor(cel(row, "valor"))
        except ValueError:
            res.avisos.append(f"Linha {n}: valor {cel(row, 'valor')!r} não reconhecido; deixado em branco.")
            valor = ""
        doc_bruto = cel(row, "documento")
        if isinstance(doc_bruto, int | float):  # Excel converte CNPJ em número e perde zeros à esquerda
            doc_bruto = f"{int(doc_bruto):014d}" if int(doc_bruto) > 99999999999 else f"{int(doc_bruto):011d}"
        res.linhas.append(
            LinhaImportada(
                nome=_sem_apostrofo(str(cel(row, "nome") or "").strip()),
                documento=normalizar_documento(str(doc_bruto or "")) or "",
                valor=valor,
                descricao=_sem_apostrofo(str(cel(row, "descricao") or "").strip()),
            )
        )
    return res


def _sem_apostrofo(v: str) -> str:
    """Desfaz o apóstrofo que o exportar_csv põe antes de =, +, - e @ (reimportar devolve o texto original)."""
    return v[1:] if v.startswith("'") and v[1:].startswith(_INICIO_FORMULA) else v


MAX_LINHAS = 10_000
MAX_BYTES_IMPORTACAO = 10_000_000


def _limitar(linhas) -> list[list]:
    saida = []
    for r in linhas:
        if len(saida) > MAX_LINHAS:  # +1 do cabeçalho
            raise ImportacaoError(
                f"A planilha passa de {MAX_LINHAS:,} linhas. Divida em arquivos menores.".replace(",", ".")
            )
        saida.append(list(r))
    return saida


def importar(nome_arquivo: str, conteudo: bytes) -> ResultadoImportacao:
    if len(conteudo) > MAX_BYTES_IMPORTACAO:
        raise ImportacaoError("Arquivo grande demais (máx. 10 MB).")
    ext = nome_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_arquivo else ""
    if ext in ("xlsx", "xlsm"):
        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl lança vários tipos para arquivo inválido
            raise ImportacaoError(f"Não foi possível ler a planilha: {exc}") from exc
        ws = wb.worksheets[0]
        return _linhas_para_resultado(_limitar(ws.iter_rows(values_only=True)))
    if ext in ("csv", "txt"):
        for enc in ("utf-8-sig", "cp1252"):
            try:
                texto = conteudo.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = _PontoEVirgula
        return _linhas_para_resultado(_limitar(csv.reader(io.StringIO(texto), dialeto)))
    raise ImportacaoError("Formato não suportado. Use .xlsx ou .csv.")


_INICIO_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _texto_csv(v: object) -> str:
    """Texto que o Excel/LibreOffice não executa como fórmula ao abrir o CSV."""
    v = str(v or "")
    return "'" + v if v.startswith(_INICIO_FORMULA) else v


def exportar_xlsx(linhas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0098D7")
        c.alignment = Alignment(vertical="center")
    for ln in linhas:
        valor = Decimal(ln["valor"]) if ln.get("valor") else None
        ws.append([ln.get("nome", ""), ln.get("documento", ""), valor, ln.get("descricao", "")])
        for col in (1, 4):  # texto digitado nunca vira fórmula na planilha exportada
            c = ws.cell(row=ws.max_row, column=col)
            if c.data_type == "f":
                c.data_type = "s"
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        row[0].number_format = "@"
        row[1].number_format = '"R$" #,##0.00'
    for i, larg in enumerate((45, 20, 14, 50), start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_csv(linhas: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(COLUNAS)
    for ln in linhas:
        valor = (ln.get("valor") or "").replace(".", ",")
        w.writerow([_texto_csv(ln.get("nome")), ln.get("documento", ""), valor, _texto_csv(ln.get("descricao"))])
    return buf.getvalue().encode("utf-8-sig")
