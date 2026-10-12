"""Regressões da auditoria de segurança de 2026-10-11."""

from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook

from emissor.db import Configuracao
from emissor.nfse.constants import Ambiente
from emissor.services import clientes, tabela


def test_exportacao_nao_gera_formula():
    linhas = [
        {
            "nome": '=HYPERLINK("http://x.invalid";"a")',
            "documento": "11222333000181",
            "valor": "10",
            "descricao": "@SUM(1)",
        }
    ]
    ws = load_workbook(io.BytesIO(tabela.exportar_xlsx(linhas))).active
    assert ws["A2"].data_type == "s" and ws["D2"].data_type == "s"
    csv = tabela.exportar_csv(linhas).decode("utf-8-sig").splitlines()[1]
    assert csv.startswith("\"'=HYPERLINK") and csv.endswith(";'@SUM(1)")
    # Reimportar o CSV devolve o texto original.
    res = tabela.importar("t.csv", tabela.exportar_csv(linhas))
    assert res.linhas[0].nome == linhas[0]["nome"] and res.linhas[0].descricao == "@SUM(1)"


def test_importacao_limita_linhas(monkeypatch):
    monkeypatch.setattr(tabela, "MAX_LINHAS", 3)
    csv = "Nome;CNPJ/CPF;Valor;Descrição\n" + "A;11222333000181;1;x\n" * 4
    try:
        tabela.importar("t.csv", csv.encode())
    except tabela.ImportacaoError as exc:
        assert "linhas" in str(exc)
    else:
        raise AssertionError("deveria recusar")


def test_zip_de_xmls_tem_limite_somado(monkeypatch):
    monkeypatch.setattr(clientes, "MAX_ZIP_DESCOMPACTADO", 2_500_000)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("a.xml", b"<a>" + b" " * 1_000_000 + b"</a>")
    um_zip = buf.getvalue()
    res = clientes.ResultadoImportacaoXml()
    saida = clientes._expandir_arquivos([(f"z{i}.zip", um_zip) for i in range(4)], res)
    assert len(saida) == 2 and len(res.ignorados) == 2


def test_ndps_usa_o_ambiente_do_lote():
    cfg = Configuracao(ambiente=Ambiente.HOMOLOGACAO.value, proximo_ndps_homologacao=5, proximo_ndps_producao=40)
    assert cfg.proximo_ndps(Ambiente.PRODUCAO) == 40
    cfg.avancar_ndps(40, Ambiente.PRODUCAO)
    assert (cfg.proximo_ndps_producao, cfg.proximo_ndps_homologacao) == (41, 5)
