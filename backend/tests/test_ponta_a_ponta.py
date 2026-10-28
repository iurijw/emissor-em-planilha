"""Testes ponta a ponta contra a Sefin/ADN simulados (``sefin_simulada.py``), sem certificado real.

Passam pelo sistema inteiro como em produção (API → fila → montagem/assinatura da DPS → "Sefin" → XML/PDF →
eventos), trocando só a rede: todo ``SefinClient`` criado sem transporte fala com a simulação, que confere a DPS
no XSD oficial e a assinatura com o certificado da conexão, e devolve uma NFS-e montada a partir dela.
O último teste chama **todas** as rotas da API e falha se aparecer uma rota nova sem teste ou um erro 500.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterator

import httpx
import pytest
from fastapi.testclient import TestClient
from lxml import etree
from openpyxl import load_workbook

from emissor import auth
from emissor.api.app import create_app
from emissor.db import agora
from emissor.nfse.certificate import Certificado
from emissor.nfse.client import SefinClient, gzip_b64, ungzip_b64
from emissor.nfse.constants import NS_NFSE
from emissor.nfse.documentos import formatar_documento
from emissor.nfse.signer import verificar
from emissor.nfse.xml_reader import NFSeDoc
from emissor.nfse.xsd import validar_dps, validar_evento, validar_nfse
from emissor.services import clientes as clientes_service
from emissor.services import emissao as emissao_service
from emissor.services import eventos as eventos_service
from emissor.services import parametros, variaveis

from .conftest import ADMIN_EMAIL, ADMIN_SENHA, CNPJ_PRESTADOR, SENHA_PFX, _gerar_pfx, entrar_como_admin
from .sefin_simulada import SefinSimulada, brasilapi_simulada, conferir_assinatura_sefin

N = f"{{{NS_NFSE}}}"
CNPJ_SEFIN = "00394460005887"  # fictício: só assina as NFS-e e os eventos da simulação

TOMADOR_CNPJ = "11222333000181"
TOMADOR_CPF = "52998224725"


@pytest.fixture(scope="session")
def cert_sefin() -> Certificado:
    return Certificado(_gerar_pfx(CNPJ_SEFIN, "sefin"), "sefin")


@pytest.fixture
def sim(cert_sefin, monkeypatch) -> SefinSimulada:
    """Liga a simulação: todo ``SefinClient`` sem transporte (fila, verificar, testar conexão...) fala com ela."""
    s = SefinSimulada(cert_sefin)
    monkeypatch.setattr(parametros, "ATIVA", True)  # conferência na prefeitura pelo ADN simulado
    init_original = SefinClient.__init__

    def init(self, ambiente, certificado, *, transport=None, **kw):
        init_original(self, ambiente, certificado, transport=transport or s.transport(certificado), **kw)

    monkeypatch.setattr(SefinClient, "__init__", init)
    consultar_original = clientes_service.consultar_cnpj
    monkeypatch.setattr(
        clientes_service,
        "consultar_cnpj",
        lambda doc, http=None: consultar_original(
            doc, http or httpx.Client(transport=httpx.MockTransport(brasilapi_simulada))
        ),
    )
    return s


class ClienteRegistrado(TestClient):
    """TestClient que anota cada rota chamada (``MÉTODO /caminho/{param}``) e falha em qualquer erro 500."""

    def __init__(self, *a, rotas: dict[tuple[str, re.Pattern], str], chamadas: set[str], **kw):
        super().__init__(*a, **kw)
        self._rotas = rotas
        self._chamadas = chamadas

    def request(self, method, url, *a, **kw):  # type: ignore[override]
        resp = super().request(method, url, *a, **kw)
        caminho = httpx.URL(str(url)).path
        for (metodo, padrao), modelo in self._rotas.items():
            if metodo == method.upper() and padrao.fullmatch(caminho):
                self._chamadas.add(f"{metodo} {modelo}")
                break
        assert resp.status_code < 500, f"{method} {url} → {resp.status_code}: {resp.text[:500]}"
        return resp


def _rotas_da_api(app) -> dict[tuple[str, re.Pattern], str]:
    rotas = {}
    for caminho, metodos in app.openapi()["paths"].items():
        if not caminho.startswith("/api/"):
            continue
        padrao = re.compile(re.sub(r"\{[^/]+\}", "[^/]+", caminho))
        for metodo in metodos:
            rotas[(metodo.upper(), padrao)] = caminho
    # Rotas fixas antes das com parâmetro (ex.: /clientes/importar-xmls antes de /clientes/{documento}).
    return dict(sorted(rotas.items(), key=lambda kv: kv[1].count("{")))


@pytest.fixture
def chamadas() -> set[str]:
    return set()


@pytest.fixture
def api(sim, chamadas) -> Iterator[ClienteRegistrado]:
    w = emissao_service.EmissionWorker(dormir=lambda s: None)  # fábrica padrão = SefinClient (simulado)
    emissao_service.definir_worker(w)
    eventos_service.definir_sincronizador(eventos_service.Sincronizador(dormir=lambda s: None))
    app = create_app(iniciar_worker=False)
    with ClienteRegistrado(app, rotas=_rotas_da_api(app), chamadas=chamadas) as c:
        c.worker = w
        c.app_rotas = _rotas_da_api(app)
        entrar_como_admin(c)
        yield c


# --- ajudantes --------------------------------------------------------------------------------------------


def _onboarding(api, pfx_bytes, nfse_exemplo_bytes) -> dict:
    r = api.post("/api/certificado", files={"arquivo": ("c.pfx", pfx_bytes)}, data={"senha": SENHA_PFX})
    assert r.status_code == 200, r.text
    assert api.get("/api/onboarding/padroes").status_code == 200
    dados = api.post("/api/onboarding/xmls", files=[("arquivos", ("n.xml", nfse_exemplo_bytes))]).json()
    r = api.post(
        "/api/onboarding/concluir",
        json={"fiscal": dados["fiscal"], "ambiente": "2", "serie": 1, "linhas": dados["linhas_sugeridas"]},
    )
    assert r.status_code == 200 and r.json()["pronto_para_emitir"], r.text
    return dados


def _linhas(api, linhas: list[dict], aba_id: int | None = None) -> list[int]:
    params = {"aba_id": aba_id} if aba_id else {}
    r = api.put("/api/tabela", params=params, json={"linhas": linhas})
    assert r.status_code == 200, r.text
    return [ln["id"] for ln in r.json()["linhas"]]


def _emitir(api, linha_ids: list[int]) -> dict:
    r = api.post("/api/lotes", json={"linha_ids": linha_ids})
    assert r.status_code == 200, r.text
    api.worker.processar_lote(r.json()["id"])
    return api.get(f"/api/lotes/{r.json()['id']}").json()


def _conferir_nota(api, sim: SefinSimulada, emissao: dict, *, documento: str, valor: str) -> NFSeDoc:
    """A nota guardada é a NFS-e que a Sefin devolveu, válida no XSD, e corresponde ao que foi pedido."""
    assert emissao["status"] == "autorizada", emissao.get("erro")
    xml = api.get(f"/api/emissoes/{emissao['id']}/xml").content
    nota = sim.notas[emissao["chave_acesso"]]
    assert xml == nota.xml
    assert validar_nfse(xml) == []
    conferir_assinatura_sefin(xml, sim.cert_sefin)
    doc = NFSeDoc.from_bytes(xml)
    dps = etree.fromstring(xml).find(f"{N}infNFSe/{N}DPS")
    assert validar_dps(dps) == []
    # A DPS embutida continua com a assinatura do prestador (a mesma conferida pela simulação ao receber).
    assert verificar(etree.tostring(dps), nota.cert_prestador_pem)
    assert doc.chave == emissao["chave_acesso"]
    assert (doc.dps("toma/CNPJ") or doc.dps("toma/CPF")) == documento
    assert doc.dps("valores/vServPrest/vServ") == valor
    assert doc.dps("prest/CNPJ") == CNPJ_PRESTADOR
    pdf = api.get(f"/api/emissoes/{emissao['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    return doc


# --- fluxos -----------------------------------------------------------------------------------------------


def test_simulacao_recusa_dps_adulterada(sim, certificado, dps_input):
    """Garante que a simulação confere de verdade (senão os testes abaixo passariam sem provar nada)."""
    from emissor.nfse.emissao import preparar_dps

    _, dps = preparar_dps(dps_input, certificado)
    with SefinClient(dps_input.ambiente, certificado) as c:

        def enviar(xml: bytes) -> httpx.Response:
            return c._http.post(f"{dps_input.ambiente.sefin_url}/nfse", json={"dpsXmlGZipB64": gzip_b64(xml)})

        adulterada = dps.replace(b"<vServ>450.00</vServ>", b"<vServ>4500.00</vServ>")
        assert enviar(adulterada).json()["erros"][0]["Codigo"] == "SIM-ASSIN"
        assert enviar(dps.replace(b"<serie>1</serie>", b"<serie>A</serie>")).json()["erros"][0]["Codigo"] == "SIM-XSD"
        outro = Certificado(_gerar_pfx("11222333000181", "x"), "x")
        with SefinClient(dps_input.ambiente, outro) as c2:
            r = c2._http.post(f"{dps_input.ambiente.sefin_url}/nfse", json={"dpsXmlGZipB64": gzip_b64(dps)})
            assert r.json()["erros"][0]["Codigo"] == "SIM-ASSIN"  # assinada por outro certificado
        r = enviar(dps)
        assert r.status_code == 201 and validar_nfse(ungzip_b64(r.json()["nfseXmlGZipB64"])) == []
        assert enviar(dps).json()["erros"][0]["Codigo"] == "E0014"


def test_lote_com_ibscbs_desligado_e_ligado(api, sim, pfx_bytes, nfse_exemplo_bytes):
    dados = _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    ids = _linhas(
        api,
        [
            {"nome": "CLIENTE EXEMPLO LTDA", "documento": TOMADOR_CNPJ, "valor": "490", "descricao": "honorarios"},
            {"nome": "PESSOA FISICA", "documento": TOMADOR_CPF, "valor": "1.234,56", "descricao": "consultoria"},
        ],
    )
    lote = _emitir(api, ids)
    assert lote["status"] == "concluido" and lote["contagem"] == {"autorizada": 2}, lote
    assert sim.rejeicoes == []
    d1 = _conferir_nota(api, sim, lote["emissoes"][0], documento=TOMADOR_CNPJ, valor="490.00")
    d2 = _conferir_nota(api, sim, lote["emissoes"][1], documento=TOMADOR_CPF, valor="1234.56")
    assert d1.el("IBSCBS") is None and d1.el("DPS/infDPS/IBSCBS") is None
    assert [d1.dps("nDPS"), d2.dps("nDPS")] == ["1", "2"]
    assert d1.dps("tpAmb") == "2" and d1.dps("serie") == "1"

    fiscal = dados["fiscal"] | {"ibscbs": True, "ibscbs_c_class_trib": "200052"}
    assert api.put("/api/config", json={"fiscal": fiscal}).status_code == 200
    lote = _emitir(api, ids[:1])
    assert lote["contagem"] == {"autorizada": 1}, lote
    d3 = _conferir_nota(api, sim, lote["emissoes"][0], documento=TOMADOR_CNPJ, valor="490.00")
    # homologação = leiaute da NT 009: finNFSe em infDPS e CST/cClassTrib direto em IBSCBS/valores/trib
    assert d3.dps("finNFSe") == "0" and d3.dps("IBSCBS/finNFSe") is None
    assert d3.dps("IBSCBS/valores/trib/CST") == "200"
    assert d3.dps("IBSCBS/valores/trib/cClassTrib") == "200052"
    assert d3.t("IBSCBS/totCIBS/gCBS/vCBS") == "4.32"  # (490 - ISS 9,80) x 0,9% calculado pela "Sefin"
    assert d3.dps("nDPS") == "3"


def test_aba_com_tributacao_e_nota_avulsa(api, sim, pfx_bytes, nfse_exemplo_bytes):
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    t = api.post(
        "/api/tributacoes",
        json={"nome": "ISS retido", "dados": {"c_trib_nac": "010101", "tp_ret_issqn": "2", "ibscbs": True}},
    ).json()
    aba = api.post("/api/abas", json={"nome": "Retenção", "tributacao_id": t["id"]}).json()
    ids = _linhas(api, [{"nome": "X", "documento": TOMADOR_CNPJ, "valor": "1000", "descricao": "servico"}], aba["id"])
    lote = _emitir(api, ids)
    doc = _conferir_nota(api, sim, lote["emissoes"][0], documento=TOMADOR_CNPJ, valor="1000.00")
    assert doc.dps("serv/cServ/cTribNac") == "010101" and doc.dps("valores/trib/tribMun/tpRetISSQN") == "2"
    assert doc.t("valores/vTotalRet") == "20.00" and doc.t("valores/vLiq") == "980.00"
    assert doc.dps("IBSCBS/valores/trib/cClassTrib") == "000001"

    nota = {"nome": "PESSOA FISICA", "documento": "529.982.247-25", "valor": "77,10", "descricao": "parecer avulso"}
    for extra, c_trib in (({}, "171901"), ({"tributacao": {"c_trib_nac": "020101"}}, "020101")):
        r = api.post("/api/avulsa", json=nota | extra)
        assert r.status_code == 200, r.text
        api.worker.processar_lote(r.json()["id"])
        e = api.get(f"/api/lotes/{r.json()['id']}").json()["emissoes"][0]
        doc = _conferir_nota(api, sim, e, documento=TOMADOR_CPF, valor="77.10")
        assert doc.dps("serv/cServ/cTribNac") == c_trib and doc.el("DPS/infDPS/IBSCBS") is None
    assert len(sim.notas) == 3


def _regime(api, **campos) -> None:
    fiscal = api.get("/api/config").json()["fiscal"] | campos
    r = api.put("/api/config", json={"fiscal": fiscal})
    assert r.status_code == 200, r.text


def test_iss_retido_no_simples_e_tributos_federais(api, sim, pfx_bytes, nfse_exemplo_bytes):
    """ME/EPP com ISS pelo Simples e ISS retido: a alíquota é obrigatória e é pedida na emissão (E0621)."""
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    _regime(api, reg_ap_trib_sn="1")
    dados = {
        "tp_ret_issqn": "2",
        "cst_pis_cofins": "01",
        "p_aliq_pis": "0,65",
        "p_aliq_cofins": "3",
        "tp_ret_pis_cofins": "3",
        "p_ret_csll": "1",
        "p_ret_irrf": "1,5",
    }
    t = api.post("/api/tributacoes", json={"nome": "Retido SN", "dados": dados}).json()
    assert t["dados"]["p_aliq_pis"] == "0.65" and t["dados"]["p_aliq_iss"] is None
    aba = api.post("/api/abas", json={"nome": "Retenção", "tributacao_id": t["id"]}).json()
    ids = _linhas(api, [{"nome": "X", "documento": TOMADOR_CNPJ, "valor": "1000", "descricao": "servico"}], aba["id"])

    v = api.post("/api/tabela/validar", json={"linha_ids": ids}).json()
    assert v["linhas"][0]["erros"] == []
    assert v["aliquotas_iss"] == [
        {
            "origem": "tributacao",
            "tributacao_id": t["id"],
            "nome": "Retido SN",
            "p_aliq_iss": None,
            "atualizada_em": None,
            "linhas": 1,
        }
    ]
    r = api.post("/api/lotes", json={"linha_ids": ids})
    assert r.status_code == 400 and "alíquota do ISS" in r.json()["mensagem"]

    assert (
        api.post("/api/tributacoes/aliquota-iss", json={"tributacao_id": t["id"], "p_aliq_iss": "9"}).status_code == 400
    )
    r = api.post("/api/tributacoes/aliquota-iss", json={"tributacao_id": t["id"], "p_aliq_iss": "2,5"})
    assert r.status_code == 200, r.text
    item = api.post("/api/tabela/validar", json={"linha_ids": ids}).json()["aliquotas_iss"][0]
    assert item["p_aliq_iss"] == "2.50" and item["atualizada_em"]

    doc = _conferir_nota(api, sim, _emitir(api, ids)["emissoes"][0], documento=TOMADOR_CNPJ, valor="1000.00")
    tf = "valores/trib/tribFed"
    assert doc.dps("valores/trib/tribMun/pAliq") == "2.50"
    assert (doc.dps(f"{tf}/piscofins/vPis"), doc.dps(f"{tf}/piscofins/vCofins")) == ("6.50", "30.00")
    assert (doc.dps(f"{tf}/vRetIRRF"), doc.dps(f"{tf}/vRetCSLL"), doc.dps(f"{tf}/vRetCP")) == ("15.00", "10.00", None)
    assert doc.t("valores/vTotalRet") == "50.00" and doc.t("valores/vLiq") == "913.50"

    # Retenções que passam do valor da nota bloqueiam a linha antes do envio.
    api.put(f"/api/tributacoes/{t['id']}", json={"nome": "Retido SN", "dados": t["dados"] | {"p_ret_cp": "99"}})
    erros = api.post("/api/tabela/validar", json={"linha_ids": ids}).json()["linhas"][0]["erros"]
    assert any("passam do valor" in e for e in erros)


def test_padroes_com_iss_retido_e_prestador_nao_optante(api, sim, pfx_bytes, nfse_exemplo_bytes):
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    _regime(api, reg_ap_trib_sn="1", tp_ret_issqn="2")
    nota = {"nome": "X", "documento": TOMADOR_CNPJ, "valor": "200", "descricao": "avulsa"}
    v = api.post("/api/avulsa/validar", json=nota).json()
    assert [i["origem"] for i in v["aliquotas_iss"]] == ["padrao"]
    assert api.post("/api/tributacoes/aliquota-iss", json={"p_aliq_iss": "3.1"}).status_code == 200
    assert api.get("/api/config").json()["fiscal"]["p_aliq_iss"] == "3.10"
    # Tributação só desta nota: a alíquota vem junto com ela (nada é salvo).
    v = api.post("/api/avulsa/validar", json=nota | {"tributacao": {"p_aliq_iss": "4"}}).json()
    assert v["aliquotas_iss"][0]["origem"] == "nota" and v["aliquotas_iss"][0]["p_aliq_iss"] == "4.00"
    r = api.post("/api/avulsa", json=nota | {"tributacao": {"p_aliq_iss": "4"}})
    api.worker.processar_lote(r.json()["id"])
    doc = _conferir_nota(
        api, sim, api.get(f"/api/lotes/{r.json()['id']}").json()["emissoes"][0], documento=TOMADOR_CNPJ, valor="200.00"
    )
    assert doc.dps("valores/trib/tribMun/pAliq") == "4.00" and doc.t("valores/vLiq") == "192.00"

    # Não optante: alíquota do ISS não vai (a do município vale) e os totais aproximados vão por esfera (E0713).
    _regime(
        api,
        op_simp_nac="1",
        reg_ap_trib_sn=None,
        tp_ret_issqn="1",
        p_tot_trib_sn=None,
        p_tot_trib_fed="13,45",
        p_tot_trib_mun="2",
    )
    r = api.post("/api/avulsa", json=nota)
    api.worker.processar_lote(r.json()["id"])
    doc = _conferir_nota(
        api, sim, api.get(f"/api/lotes/{r.json()['id']}").json()["emissoes"][0], documento=TOMADOR_CNPJ, valor="200.00"
    )
    tot = "valores/trib/totTrib/pTotTrib"
    assert (doc.dps(f"{tot}/pTotTribFed"), doc.dps(f"{tot}/pTotTribEst"), doc.dps(f"{tot}/pTotTribMun")) == (
        "13.45",
        "0.00",
        "2.00",
    )
    assert doc.dps("valores/trib/tribMun/pAliq") is None and doc.dps("valores/trib/totTrib/indTotTrib") is None


def test_cancelamento_e_situacao_pelo_adn(api, sim, pfx_bytes, nfse_exemplo_bytes):
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    ids = _linhas(
        api,
        [
            {"nome": "A", "documento": TOMADOR_CNPJ, "valor": "10", "descricao": "a"},
            {"nome": "B", "documento": TOMADOR_CPF, "valor": "20", "descricao": "b"},
        ],
    )
    e1, e2 = _emitir(api, ids)["emissoes"]

    r = api.post(f"/api/emissoes/{e1['id']}/cancelar", json={"c_motivo": "1", "x_motivo": "Valor digitado errado"})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "cancelada"
    evento = sim.notas[e1["chave_acesso"]].eventos["101101"]
    assert validar_evento(evento) == []
    conferir_assinatura_sefin(evento, sim.cert_sefin)
    ev = r.json()["eventos"][0]
    assert api.get(f"/api/emissoes/{e1['id']}/eventos/{ev['id']}/xml").content == evento
    r = api.post(f"/api/emissoes/{e1['id']}/cancelar", json={"c_motivo": "1", "x_motivo": "Valor digitado errado"})
    assert r.status_code == 400  # já cancelada: nem chega à Sefin
    assert len(sim.pedidos_recebidos) == 1

    sim.cancelar_no_portal(e2["chave_acesso"])
    eventos_service.sincronizador().executar()
    assert api.get(f"/api/emissoes/{e2['id']}").json()["status"] == "cancelada"
    assert api.get(f"/api/emissoes/{e2['id']}").json()["eventos"][0]["origem"] == "ambiente_nacional"
    assert api.get("/api/emissoes", params={"status": "cancelada"}).json()["total"] == 2


def test_rejeicao_queda_e_certificado_recusado(api, sim, pfx_bytes, nfse_exemplo_bytes):
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    sim.rejeitar_tomador[TOMADOR_CPF] = ("E0312", "Tomador não pode ser identificado (simulação)")
    sim.cair_apos_processar = 1  # a 1ª nota é gerada, mas a resposta se perde (timeout)
    ids = _linhas(
        api,
        [
            {"nome": "A", "documento": TOMADOR_CNPJ, "valor": "10", "descricao": "a"},
            {"nome": "B", "documento": TOMADOR_CPF, "valor": "20", "descricao": "b"},
            {"nome": "C", "documento": TOMADOR_CNPJ, "valor": "30", "descricao": "c"},
        ],
    )
    lote = _emitir(api, ids)
    assert [e["status"] for e in lote["emissoes"]] == ["autorizada", "rejeitada", "autorizada"], lote
    # Sem duplicidade: a nota "perdida" foi recuperada pela consulta da DPS, não reenviada como outra nota.
    assert len(sim.notas) == 2
    _conferir_nota(api, sim, lote["emissoes"][0], documento=TOMADOR_CNPJ, valor="10.00")
    assert lote["emissoes"][1]["erro"]["mensagens"][0]["codigo"] == "E0312"

    sim.certificados_recusados.add(CNPJ_PRESTADOR)
    r = api.post("/api/config/testar-conexao").json()
    assert r["ok"] is False and r["erro"]["tipo"] == "autenticacao"
    lote = _emitir(api, ids)
    assert lote["status"] == "interrompido"
    assert [e["status"] for e in lote["emissoes"]] == ["erro", "nao_enviada", "nao_enviada"]


def test_producao_usa_outro_endereco_e_outra_numeracao(api, sim, pfx_bytes, nfse_exemplo_bytes):
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    ids = [ln["id"] for ln in api.get("/api/tabela").json()["linhas"]]
    assert _emitir(api, ids)["contagem"] == {"autorizada": 1}
    r = api.put("/api/config", json={"ambiente": "1", "proximo_ndps_producao": 500})
    assert r.status_code == 200, r.text
    e = _emitir(api, ids)["emissoes"][0]
    doc = _conferir_nota(api, sim, e, documento=TOMADOR_CNPJ, valor="490.00")
    assert doc.dps("tpAmb") == "1" and doc.dps("nDPS") == "500"
    assert any(c.startswith("POST sefin.nfse.gov.br/SefinNacional/nfse") for c in sim.chamadas)


def test_numero_da_dps_ja_usado_por_outro_sistema(api, sim, pfx_bytes, nfse_exemplo_bytes):
    """E0014 com a NFS-e de outra nota: não adota a nota alheia como se fosse a emitida (bug da 0.1.0)."""
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    ids = _linhas(
        api,
        [
            {"nome": "A", "documento": TOMADOR_CNPJ, "valor": "10", "descricao": "a"},
            {"nome": "B", "documento": TOMADOR_CPF, "valor": "20", "descricao": "b"},
        ],
    )
    assert _emitir(api, ids[:1])["contagem"] == {"autorizada": 1}
    # Outro sistema com o mesmo CNPJ e série (ou o contador voltou): o próximo número já foi usado.
    assert api.put("/api/config", json={"proximo_ndps_homologacao": 1}).status_code == 200
    lote = _emitir(api, [ids[1], ids[0]])
    assert lote["status"] == "interrompido" and "numeração" in lote["motivo_interrupcao"]
    rej, nao = lote["emissoes"]
    assert rej["status"] == "rejeitada" and rej["chave_acesso"] is None, rej
    msg = rej["erro"]["mensagens"][0]
    assert msg["codigo"] == "E0014" and "nº 1 já gerou a NFS-e" in msg["descricao"] and msg["dica"]
    assert nao["status"] == "nao_enviada"
    assert len(sim.notas) == 1  # nada foi emitido nem reaproveitado
    assert api.get("/api/emissoes", params={"status": "autorizada"}).json()["total"] == 1


def test_servico_que_a_prefeitura_nao_administra(api, sim, pfx_bytes, nfse_exemplo_bytes):
    """E0312: a conferência pelo ADN avisa antes de emitir (sem bloquear) e a rejeição traz a dica."""
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    fiscal = api.get("/api/config").json()["fiscal"]
    c_trib, mun = fiscal["c_trib_nac"], fiscal["c_mun_emissor"]

    def conferir(**kw) -> dict:
        r = api.post("/api/config/conferir-servico", json={"c_trib_nac": c_trib, **kw})
        assert r.status_code == 200, r.text
        return r.json()

    def consultas() -> int:
        return sum("/parametrizacao/" in c for c in sim.chamadas)

    r = conferir()
    assert r["situacao"] == "encontrado" and "2,00%" in r["mensagem"] and r["municipio"] == mun, r
    assert consultas() == 1
    ids = _linhas(api, [{"nome": "A", "documento": TOMADOR_CNPJ, "valor": "10", "descricao": "a"}])
    assert api.post("/api/tabela/validar", json={"linha_ids": ids}).json()["avisos"] == []
    assert consultas() == 1  # resposta guardada: a validação não consulta de novo

    sim.servicos_nao_administrados.add((mun, c_trib))
    parametros.limpar_cache()
    r = conferir()
    assert r["situacao"] == "nao_encontrado" and "E0312" in r["mensagem"], r
    assert consultas() == 3  # sem cadastro no formato de referência, tenta o alternativo (que a simulação recusa)
    val = api.post("/api/tabela/validar", json={"linha_ids": ids}).json()
    assert len(val["avisos"]) == 1 and "(tributação padrão)" in val["avisos"][0], val
    assert val["linhas"][0]["erros"] == []  # aviso não bloqueia
    av = api.post(
        "/api/avulsa/validar", json={"nome": "A", "documento": TOMADOR_CNPJ, "valor": "10", "descricao": "a"}
    ).json()
    assert len(av["avisos"]) == 1
    rej = _emitir(api, ids)["emissoes"][0]
    assert rej["status"] == "rejeitada", rej
    msg = rej["erro"]["mensagens"][0]
    assert msg["codigo"] == "E0312" and "a tela consulta a prefeitura" in msg["dica"]

    # Local da prestação em outro município (o ISS pode ser de outra prefeitura) e ADN fora do ar: sem conferir.
    assert conferir(c_loc_prestacao="3550308")["situacao"] == "nao_conferido"
    parametros.limpar_cache()
    sim.fora_do_ar = True
    r = conferir()
    assert r["situacao"] == "nao_conferido" and r["mensagem"].startswith("Não foi possível conferir agora"), r
    sim.fora_do_ar = False
    assert api.post("/api/tabela/validar", json={"linha_ids": ids}).json()["avisos"] == []  # falha guardada 5 min


# --- todas as rotas ---------------------------------------------------------------------------------------


def test_variaveis_na_descricao_e_relatorio(api, sim, pfx_bytes, nfse_exemplo_bytes):
    """A descrição sai com as variáveis trocadas pela data da emissão; o relatório soma ISS, retenções e líquido."""
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    _regime(api, reg_ap_trib_sn="1")
    t = api.post(
        "/api/tributacoes",
        json={"nome": "Retido", "dados": {"tp_ret_issqn": "2", "p_aliq_iss": "2", "p_ret_irrf": "1,5"}},
    ).json()
    aba = api.post("/api/abas", json={"nome": "Retenção", "tributacao_id": t["id"]}).json()
    texto = "Honorários de {Mês}/{ano} (ref. {mes_anterior}, {competencia_anterior})"
    ids = _linhas(
        api,
        [
            {"nome": "X", "documento": TOMADOR_CNPJ, "valor": "1000", "descricao": texto},
            {"nome": "Y", "documento": TOMADOR_CPF, "valor": "200", "descricao": "Serviço de {mes_atual}"},
        ],
        aba["id"],
    )
    v = api.post("/api/tabela/validar", json={"linha_ids": ids}).json()["linhas"]
    assert v[0]["erros"] == []
    assert any("{mes_atual}" in e for e in v[1]["erros"])
    nomes = [i["valor"] for i in api.get("/api/config/opcoes").json()["variaveis_descricao"]]
    assert "{mes}" in nomes and "{competencia_anterior}" in nomes

    ids = _linhas(
        api, [{"nome": "Y", "documento": TOMADOR_CNPJ, "valor": "200", "descricao": "Serviço {MES}"}], aba["id"]
    )
    ids += _linhas(api, [{"nome": "Z", "documento": TOMADOR_CPF, "valor": "300", "descricao": texto}])
    lote = _emitir(api, ids)
    assert lote["contagem"] == {"autorizada": 2}, lote
    e = lote["emissoes"][0]
    hoje = agora().date()
    esperado = variaveis.substituir(texto, hoje)
    assert "{" not in esperado and variaveis.MESES[hoje.month - 1].capitalize() in esperado
    doc = _conferir_nota(api, sim, e, documento=TOMADOR_CNPJ, valor="200.00")
    assert doc.dps("serv/cServ/xDescServ") == f"Serviço {variaveis.MESES[hoje.month - 1].upper()}"
    assert api.get(f"/api/emissoes/{e['id']}").json()["descricao"] == doc.dps("serv/cServ/xDescServ")
    outra = lote["emissoes"][1]
    assert api.get(f"/api/emissoes/{outra['id']}").json()["descricao"] == esperado
    # A tabela continua com as variáveis, para o mês que vem.
    linha = api.get("/api/tabela", params={"aba_id": aba["id"]}).json()["linhas"][0]
    assert linha["descricao"] == "Serviço {MES}"

    # Planilha das notas: retenção e líquido lidos do XML; a cancelada aparece com a situação.
    api.post(f"/api/emissoes/{outra['id']}/cancelar", json={"c_motivo": "2", "x_motivo": "Serviço não prestado"})
    r = api.get("/api/emissoes/relatorio")
    assert r.status_code == 200, r.text
    assert 'filename="notas-nfse-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Notas"]
    notas = list(wb["Notas"].iter_rows(values_only=True))
    cab = notas[0]
    linhas_rel = [dict(zip(cab, ln, strict=True)) for ln in notas[1:]]
    assert len(linhas_rel) == 2
    retida = next(ln for ln in linhas_rel if ln["CNPJ/CPF"] == formatar_documento(TOMADOR_CNPJ))
    assert retida["Situação"] == "Autorizada"
    assert retida["IRRF retido"] == 3 and retida["ISS retido pelo tomador"] == 4
    assert retida["Total retido"] == 7 and retida["Valor líquido"] == 193
    cancelada = next(ln for ln in linhas_rel if ln["CNPJ/CPF"] == formatar_documento(TOMADOR_CPF))
    assert cancelada["Situação"] == "Cancelada" and cancelada["Valor do serviço"] == 300
    assert api.get("/api/emissoes/relatorio", params={"status": "rejeitada"}).status_code == 400
    assert api.get("/api/emissoes/relatorio", params={"formato": "pdf"}).status_code == 400


def test_todas_as_rotas(api, sim, chamadas, pfx_bytes, nfse_exemplo_bytes):
    """Chama cada rota da API pelo menos uma vez num uso completo; rota nova sem teste faz este teste falhar."""

    def ok(r: httpx.Response, status: int = 200) -> httpx.Response:
        assert r.status_code == status, f"{r.request.method} {r.request.url} → {r.status_code}: {r.text[:500]}"
        return r

    # Público e sessão
    assert ok(api.get("/api/saude")).json()["administrador_criado"] is True
    assert ok(api.get("/api/auth/sessao")).json()["usuario"]["email"] == ADMIN_EMAIL
    ok(api.get("/api/auth/conta"))
    ok(api.post("/api/auth/sair"))
    ok(api.get("/api/status"), 401)
    ok(api.post("/api/auth/entrar", json={"email": ADMIN_EMAIL, "senha": ADMIN_SENHA}))
    ok(api.post("/api/auth/primeiro-acesso", json={"email": "x@x.com", "senha": "outra-senha-123"}), 409)
    ok(api.post("/api/auth/esqueci", json={"email": ADMIN_EMAIL}), 404)  # sem e-mail configurado

    # Conta: senha e 2FA
    nova = ADMIN_SENHA + "-nova"
    ok(api.post("/api/auth/senha", json={"senha_atual": ADMIN_SENHA, "nova_senha": nova}))
    segredo = ok(api.post("/api/auth/2fa/iniciar", json={"senha": nova})).json()["segredo"]
    ok(api.post("/api/auth/2fa/ativar", json={"codigo": auth.totp_agora(segredo)}))
    ok(api.post("/api/auth/2fa/desativar", json={"senha": nova, "codigo": "000000"}), 400)

    # Usuários e convite
    convite = ok(api.post("/api/usuarios", json={"email": "op@exemplo.com.br", "nome": "Op"})).json()
    token = convite["link"].split("token=")[1]
    usuarios = ok(api.get("/api/usuarios")).json()
    op_id = next(u["id"] for u in usuarios if u["email"] == "op@exemplo.com.br")
    ok(api.put(f"/api/usuarios/{op_id}", json={"nome": "Operador"}))
    with TestClient(api.app) as op:
        assert op.post("/api/auth/convite/ver", json={"token": token}).json()["email"] == "op@exemplo.com.br"
        assert op.post("/api/auth/convite", json={"token": token, "senha": "operador-senha-1"}).status_code == 200
        assert op.put("/api/config", json={}).status_code == 403  # operador não mexe em configurações
    for caminho, corpo in (
        ("/api/auth/convite/ver", {"token": "x"}),
        ("/api/auth/convite", {"token": "x", "senha": "s" * 12}),
    ):
        assert api.post(caminho, json=corpo).status_code in (400, 404, 410)
    ok(api.post(f"/api/usuarios/{op_id}/redefinir", json={}))
    ok(api.get("/api/auditoria"))
    # Assinatura da hospedagem: sem o painel (instalação própria) não há nada a confirmar.
    assert ok(api.get("/api/assinatura")).json() == {"disponivel": False, "pode_confirmar": True}
    ok(api.post("/api/assinatura/confirmar"), 404)
    ok(api.post("/api/assinatura/recusar"), 404)

    # Onboarding e configurações
    _onboarding(api, pfx_bytes, nfse_exemplo_bytes)
    ok(api.get("/api/status"))
    cfg = ok(api.get("/api/config")).json()
    ok(api.put("/api/config", json={"fiscal": cfg["fiscal"]}))
    ok(api.get("/api/config/opcoes"))
    assert "4205407" in [m["codigo"] for m in ok(api.get("/api/municipios", params={"q": "florian"})).json()]
    assert ok(api.get("/api/servicos", params={"q": "contabilidade"})).json()[0]["codigo"] == "171901"
    assert ok(api.get("/api/servicos/17.19.01")).json()["nbs"][0]["codigo"] == "113022100"
    ok(api.get("/api/servicos/999999"), 404)
    assert (
        ok(api.post("/api/config/conferir-servico", json={"c_trib_nac": "171901"})).json()["situacao"] == "encontrado"
    )
    assert ok(api.post("/api/config/testar-conexao")).json()["ok"] is True
    ok(api.post("/api/logs/cliente", json={"mensagem": "teste de log do navegador"}))

    # Clientes (BrasilAPI simulada)
    ok(api.post("/api/clientes/importar-xmls", files=[("arquivos", ("n.xml", nfse_exemplo_bytes))]))
    assert ok(api.get("/api/clientes", params={"q": "EXEMPLO"})).json()
    ok(api.get(f"/api/clientes/{TOMADOR_CNPJ}"))
    ok(api.post(f"/api/clientes/{TOMADOR_CNPJ}/consultar"))
    ok(api.put(f"/api/clientes/{TOMADOR_CPF}", json={"nome": "PESSOA FISICA"}))
    ok(api.post("/api/clientes/consultar-pendentes", json={"documentos": ["11444777000161"]}))
    ok(api.delete(f"/api/clientes/{TOMADOR_CPF}"))

    # Tributações, abas e tabela
    t = ok(api.post("/api/tributacoes", json={"nome": "Outra", "dados": {"c_trib_nac": "010101"}})).json()
    ok(api.put(f"/api/tributacoes/{t['id']}", json={"nome": "Outra 2", "dados": {"c_trib_nac": "010101"}}))
    ok(api.get("/api/tributacoes"))
    ok(api.post("/api/tributacoes/aliquota-iss", json={"tributacao_id": t["id"], "p_aliq_iss": "2"}))
    aba = ok(api.post("/api/abas", json={"nome": "Extra", "tributacao_id": t["id"]})).json()
    ok(api.put(f"/api/abas/{aba['id']}", json={"nome": "Extra 2", "tributacao_id": t["id"]}))
    ok(api.get("/api/abas"))
    ok(api.get("/api/tabela", params={"aba_id": aba["id"]}))
    csv = f"Nome;CNPJ;Valor;Descrição\nX;{TOMADOR_CNPJ};10,00;teste\n".encode()
    ok(api.post("/api/tabela/importar", params={"aba_id": aba["id"]}, files={"arquivo": ("t.csv", csv)}))
    assert ok(api.get("/api/tabela/exportar", params={"formato": "csv", "aba_id": aba["id"]})).content
    linhas = api.get("/api/tabela", params={"aba_id": aba["id"]}).json()["linhas"]
    linhas[0]["descricao"] = "teste editado na grade"
    ids_aba = _linhas(api, linhas, aba["id"])
    ok(api.post("/api/tabela/validar", json={"linha_ids": ids_aba}))

    # Emissão em lote, nota avulsa e lote cancelado
    lote = _emitir(api, ids_aba)
    assert lote["contagem"] == {"autorizada": 1}
    ok(api.get("/api/lotes"))
    ok(api.post("/api/avulsa/validar", json={"documento": TOMADOR_CPF, "nome": "P", "valor": "5", "descricao": "x"}))
    avulsa = ok(api.post("/api/avulsa", json={"documento": TOMADOR_CPF, "nome": "P", "valor": "5", "descricao": "x"}))
    pendente = ok(api.post("/api/lotes", json={"linha_ids": ids_aba})).json()
    ok(api.post(f"/api/lotes/{pendente['id']}/cancelar"))
    api.worker.processar_lote(pendente["id"])
    cancelado = api.get(f"/api/lotes/{pendente['id']}").json()
    assert cancelado["status"] == "cancelado" and cancelado["contagem"] == {"nao_enviada": 1}, cancelado
    api.worker.processar_lote(avulsa.json()["id"])

    # Emissões: consulta, downloads, verificar, cancelar, situação e sincronização
    itens = ok(api.get("/api/emissoes", params={"status": "autorizada"})).json()["itens"]
    assert len(itens) == 2
    eid = itens[0]["id"]
    ok(api.get(f"/api/emissoes/{eid}"))
    ok(api.get(f"/api/emissoes/{eid}/xml"))
    ok(api.get(f"/api/emissoes/{eid}/pdf"))
    ok(api.post(f"/api/emissoes/{eid}/verificar"), 400)  # só para resultado incerto
    ok(api.post(f"/api/emissoes/{eid}/situacao"))
    d = ok(api.post(f"/api/emissoes/{eid}/cancelar", json={"c_motivo": "2", "x_motivo": "Serviço não prestado"}))
    ok(api.get(f"/api/emissoes/{eid}/eventos/{d.json()['eventos'][0]['id']}/xml"))
    ok(api.post("/api/sincronizacao", params={"forcar": True}))
    if eventos_service.sincronizador()._thread:
        eventos_service.sincronizador()._thread.join(10)
    ok(api.get("/api/sincronizacao"))
    z = ok(api.post("/api/emissoes/zip", json={"ids": [i["id"] for i in itens], "conteudo": "ambos"}))
    assert len(zipfile.ZipFile(io.BytesIO(z.content)).namelist()) == 5  # 2 XML + 1 evento + 2 PDF
    rel = ok(api.get("/api/emissoes/relatorio", params={"formato": "csv"}))
    assert rel.content.decode("utf-8-sig").count("\n") == 3  # cabeçalho + 2 notas
    assert sim.rejeicoes == []

    # Limpeza e diagnóstico
    ok(api.delete(f"/api/abas/{aba['id']}"))
    ok(api.delete(f"/api/tributacoes/{t['id']}"))
    assert ok(api.get("/api/diagnostico")).content.startswith(b"PK")

    faltando = sorted({f"{m} {p}" for (m, _), p in api.app_rotas.items()} - chamadas)
    assert faltando == [], f"Rotas sem teste ponta a ponta: {faltando}"
