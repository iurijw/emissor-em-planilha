"""Gera ``src/emissor/data/servicos_nacionais.json``: a lista nacional de serviços (cTribNac) com os NBS de cada um.

Fontes oficiais (gov.br/nfse, baixar e passar os caminhos):
  --anexo-b     Anexo B (lista de serviços nacional com cTribNac e lista NBS 2.0), "Documentação atual"
  --anexo-viii  Anexo VIII (correlação item LC 116 x NBS x cIndOp x cClassTrib), página "RTC"
  --classtrib   JSON da tabela de classificação tributária do Portal da Conformidade Fácil
                (https://dfe-portal.svrs.rs.gov.br/DFE/ClassificacaoTributaria, variável ``dadosOriginais`` da página)

Uso (de ``backend/``):
  uv run --with openpyxl python scripts/servicos_nacionais.py --anexo-b b.xlsx --anexo-viii a8.xlsx \\
      --classtrib ct.json --versoes "Anexo B v1.01 (22/01/2026)|Anexo VIII v1.02 (26/10/2026)" --consulta 2026-10-27

A saída é versionada (como ``municipios_ibge.csv``); o sistema só lê o JSON. Quando o gov.br publicar versão nova de
algum dos anexos, gerar de novo.
"""

import argparse
import json
import re
from pathlib import Path

import openpyxl

DESTINO = Path(__file__).resolve().parents[1] / "src/emissor/data/servicos_nacionais.json"


def texto(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def codigo6(v) -> str:
    return str(int(v)).zfill(6) if v not in (None, "") else ""


def nbs9(v: str) -> str:
    """``1.1302.21.00`` → ``113022100`` (formato do cNBS na DPS)."""
    return re.sub(r"\D", "", v)


def main() -> None:
    ap = argparse.ArgumentParser()
    for a in ("anexo-b", "anexo-viii", "classtrib", "versoes", "consulta"):
        ap.add_argument(f"--{a}", required=True)
    args = ap.parse_args()

    # Anexo B: cTribNac = item (2) + subitem (2) + desdobro nacional (2); linhas sem desdobro são grupo/subitem.
    wb = openpyxl.load_workbook(args.anexo_b, read_only=True)
    grupos, subitens, servicos = {}, {}, []
    for r in list(wb["LISTA.SERV.NAC."].iter_rows(values_only=True))[1:]:
        if r[1] is None:
            continue
        item, sub, desd = int(r[1]), int(r[2]), int(r[3])
        if sub == 0:
            grupos[item] = texto(r[4])
        elif desd == 0:
            subitens[(item, sub)] = texto(r[4])
        else:
            servicos.append({"codigo": codigo6(r[0]), "descricao": texto(r[4]), "subitem": f"{item:02d}.{sub:02d}"})
    nbs_desc = {
        nbs9(texto(r[0])): texto(r[1])
        for r in list(wb["LISTA.NBS_v2.0"].iter_rows(values_only=True))[1:]
        if r[0] and len(nbs9(texto(r[0]))) == 9
    }

    # Anexo VIII: blocos por item da LC 116 (subitem "17.19" ou desdobro "99.01.01"); cada NBS leva os cIndOp das
    # linhas abaixo dele (sem nenhum, herda os do anterior); os cClassTrib valem para o item.
    correlacao: dict[str, dict] = {}
    atual = None
    for r in list(openpyxl.load_workbook(args.anexo_viii, read_only=True)["tabela geral"].iter_rows(values_only=True))[
        1:
    ]:
        if r[0]:
            atual = correlacao.setdefault(texto(r[0]), {"nbs": [], "c_class_trib": []})
        if atual is None:
            continue
        if texto(r[2]):
            codigo = nbs9(texto(r[2]))
            atual["nbs"].append({"codigo": codigo, "descricao": nbs_desc.get(codigo) or texto(r[3]), "c_ind_op": []})
        if r[6] and atual["nbs"] and codigo6(r[6]) not in atual["nbs"][-1]["c_ind_op"]:
            atual["nbs"][-1]["c_ind_op"].append(codigo6(r[6]))
        if r[8] and codigo6(r[8]) not in atual["c_class_trib"]:
            atual["c_class_trib"].append(codigo6(r[8]))
    for c in correlacao.values():
        anterior: list[str] = []
        for n in c["nbs"]:
            n["c_ind_op"] = n["c_ind_op"] or anterior
            anterior = n["c_ind_op"]

    sem_nbs = 0
    for s in servicos:
        item, sub = int(s["subitem"][:2]), int(s["subitem"][3:])
        s["subitem_descricao"] = subitens.get((item, sub), "")
        s["grupo_descricao"] = grupos.get(item, "")
        c = correlacao.get(f"{s['subitem']}.{s['codigo'][4:]}") or correlacao.get(s["subitem"]) or {}
        s["nbs"] = c.get("nbs", [])
        s["c_class_trib"] = c.get("c_class_trib", [])
        sem_nbs += not s["nbs"]

    # cClassTrib vigentes que valem para NFS-e (Conformidade Fácil), com CST, nome e redução: a tela lista por CST,
    # como o Emissor Nacional, e explica as sugestões do Anexo VIII.
    tabela = json.loads(Path(args.classtrib).read_text(encoding="utf-8"))
    vigentes = {
        c["CodClassTrib"]: c for cst in tabela for c in cst.get("ClassificacoesTributarias") or [] if not c.get("DthFimVig")
    }
    usados = {c for s in servicos for c in s["c_class_trib"]}
    faltam = sorted(c for c in usados if c not in vigentes)
    if faltam:
        raise SystemExit(f"cClassTrib do Anexo VIII fora da tabela vigente: {faltam}")
    classes = {
        c: {
            "cst": v["Cst"],
            "nome": texto(v["NomeReduzido"]).rstrip("."),
            "reducao_ibs": v.get("PercRedIbs") or 0,
            "reducao_cbs": v.get("PercRedCbs") or 0,
        }
        for c, v in sorted(vigentes.items())
        if v.get("IndNfse") or c in usados
    }
    csts = {t["Cst"]: texto(t["NomeCst"]) for t in tabela if any(c["cst"] == t["Cst"] for c in classes.values())}

    dados = {
        "fontes": args.versoes.split("|"),
        "consulta": args.consulta,
        "servicos": servicos,
        "classes": classes,
        "csts": csts,
    }
    DESTINO.write_text(json.dumps(dados, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(
        f"{len(servicos)} serviços ({sem_nbs} sem NBS no Anexo VIII) -> {DESTINO} ({DESTINO.stat().st_size // 1024} KB)"
    )


if __name__ == "__main__":
    main()
